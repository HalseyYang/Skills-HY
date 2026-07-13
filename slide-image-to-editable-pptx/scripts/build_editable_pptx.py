#!/usr/bin/env python3
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WIDE_W = 13.333333
WIDE_H = 7.5


def color(value, default="FFFFFF"):
    if not value or value == "transparent":
        value = default
    value = value.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def box(el, prs):
    return (
        int(prs.slide_width * float(el.get("x", 0))),
        int(prs.slide_height * float(el.get("y", 0))),
        int(prs.slide_width * float(el.get("w", 0.1))),
        int(prs.slide_height * float(el.get("h", 0.1))),
    )


def set_fill(shape, value):
    if value == "transparent":
        shape.fill.background()
    elif value:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(value)


def set_line(shape, value, width=None):
    if value == "transparent":
        shape.line.fill.background()
    elif value:
        shape.line.color.rgb = color(value)
    if width:
        shape.line.width = Pt(float(width))


def add_text(slide, prs, el):
    x, y, w, h = box(el, prs)
    shape = slide.shapes.add_textbox(x, y, w, h)
    set_fill(shape, el.get("fill"))
    set_line(shape, el.get("line"))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = str(el.get("text", ""))
    align = str(el.get("align", "left")).lower()
    p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
    run = p.runs[0] if p.runs else p.add_run()
    font = run.font
    font.name = el.get("font_face", "Arial")
    font.size = Pt(float(el.get("font_size", 18)))
    font.bold = bool(el.get("bold", False))
    font.italic = bool(el.get("italic", False))
    font.color.rgb = color(el.get("color", "#111111"), "111111")


def add_rect(slide, prs, el):
    from pptx.enum.shapes import MSO_SHAPE

    x, y, w, h = box(el, prs)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    set_fill(shape, el.get("fill", "#FFFFFF"))
    set_line(shape, el.get("line", "transparent"), el.get("line_width"))


def add_line(slide, prs, el):
    x, y, w, h = box(el, prs)
    shape = slide.shapes.add_connector(1, x, y, x + w, y + h)
    set_line(shape, el.get("line", "#111111"), el.get("width", 1))


def add_table(slide, prs, el):
    rows = el.get("rows") or [[""]]
    x, y, w, h = box(el, prs)
    table_shape = slide.shapes.add_table(len(rows), max(len(r) for r in rows), x, y, w, h)
    table = table_shape.table
    font_size = Pt(float(el.get("font_size", 12)))
    for r_idx, row in enumerate(rows):
        for c_idx in range(len(table.columns)):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(row[c_idx]) if c_idx < len(row) else ""
            fill = el.get("header_fill") if r_idx == 0 else el.get("body_fill")
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = color(fill)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = font_size
                    run.font.color.rgb = color(el.get("text_color", "#111111"), "111111")


def add_image(slide, prs, el, base_dir):
    x, y, w, h = box(el, prs)
    path = Path(el["path"])
    if not path.is_absolute():
        path = base_dir / path
    slide.shapes.add_picture(str(path), x, y, w, h)


def main():
    if len(sys.argv) != 3:
        print("Usage: build_editable_pptx.py layout.json output.pptx", file=sys.stderr)
        return 2

    layout_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    data = json.loads(layout_path.read_text(encoding="utf-8"))

    prs = Presentation()
    if data.get("slide_size", "wide") == "wide":
        prs.slide_width = Inches(WIDE_W)
        prs.slide_height = Inches(WIDE_H)

    blank = prs.slide_layouts[6]
    for slide_spec in data.get("slides", []):
        slide = prs.slides.add_slide(blank)
        bg = slide_spec.get("background")
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = color(bg)

        for el in slide_spec.get("elements", []):
            kind = el.get("type")
            if kind == "text":
                add_text(slide, prs, el)
            elif kind == "rect":
                add_rect(slide, prs, el)
            elif kind == "line":
                add_line(slide, prs, el)
            elif kind == "table":
                add_table(slide, prs, el)
            elif kind == "image":
                add_image(slide, prs, el, layout_path.parent)
            else:
                raise ValueError(f"Unsupported element type: {kind}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
