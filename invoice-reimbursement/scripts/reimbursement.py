#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
发票报销处理脚本
功能：解析 PDF 发票，自动分类（车票/住宿费/打车费），计算差旅补助，生成 Excel 汇总表
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    print("[错误] 请先安装依赖：pip install pdfplumber openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[错误] 请先安装依赖：pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# ─────────────────────────────────────────────
# 城市等级映射表
# ─────────────────────────────────────────────
CITY_LEVEL = {
    "北京": "A", "北京市": "A",
    "上海": "B", "上海市": "B",
    "广州": "B", "广州市": "B",
    "深圳": "B", "深圳市": "B",
    "杭州": "C", "杭州市": "C",
    "苏州": "C", "苏州市": "C",
    "宁波": "C", "宁波市": "C",
    "青岛": "C", "青岛市": "C",
    "昆山": "C", "昆山市": "C",
    "厦门": "C", "厦门市": "C",
    "东莞": "C", "东莞市": "C",
    "佛山": "C", "佛山市": "C",
    "哈尔滨": "C", "哈尔滨市": "C",
    "长春": "C", "长春市": "C",
    "沈阳": "C", "沈阳市": "C",
    "天津": "C", "天津市": "C",
    "石家庄": "C", "石家庄市": "C",
    "太原": "C", "太原市": "C",
    "呼和浩特": "C", "呼和浩特市": "C",
    "济南": "C", "济南市": "C",
    "郑州": "C", "郑州市": "C",
    "西安": "C", "西安市": "C",
    "兰州": "C", "兰州市": "C",
    "银川": "C", "银川市": "C",
    "西宁": "C", "西宁市": "C",
    "乌鲁木齐": "C", "乌鲁木齐市": "C",
    "拉萨": "C", "拉萨市": "C",
    "成都": "C", "成都市": "C",
    "重庆": "C", "重庆市": "C",
    "长沙": "C", "长沙市": "C",
    "武汉": "C", "武汉市": "C",
    "合肥": "C", "合肥市": "C",
    "南京": "C", "南京市": "C",
    "南昌": "C", "南昌市": "C",
    "昆明": "C", "昆明市": "C",
    "贵阳": "C", "贵阳市": "C",
    "南宁": "C", "南宁市": "C",
    "海口": "C", "海口市": "C",
    "福州": "C", "福州市": "C",
}

# 城市等级 -> 补助标准（元/天）
ALLOWANCE = {"A": 200, "B": 200, "C": 150, "D": 150, "E": 150}

# ─────────────────────────────────────────────
# 发票类型关键词
# ─────────────────────────────────────────────
CATEGORY_KEYWORDS = {
    "车票": [
        "铁路", "高铁", "动车", "火车票", "火车", "机票", "航空", "航班",
        "客车", "大巴", "行程单", "旅客", "列车", "飞机", "起飞", "降落",
        "出发", "终到", "始发", "班次", "票价", "座位", "二等座", "一等座",
        "商务座", "卧铺", "交通运输", "铁道", "城际", "地铁", "城轨"
    ],
    "住宿费": [
        "酒店", "住宿", "宾馆", "旅店", "旅馆", "客房", "房费", "住宿费",
        "夜", "间夜", "客房消费", "订房", "退房", "入住", "房型", "大床房",
        "标准间", "豪华套", "民宿", "青旅", "旅社", "招待所"
    ],
    "打车费": [
        "滴滴", "出租车", "的士", "打车", "快车", "专车", "花小猪",
        "T3出行", "首汽", "曹操", "神州", "享道", "高德打车", "哈啰打车",
        "出租", "网约车", "行程账单", "叫车", "乘车", "打表", "计程"
    ]
}

# ─────────────────────────────────────────────
# PDF 解析
# ─────────────────────────────────────────────

def extract_text_from_pdf(pdf_path: str) -> str:
    """提取 PDF 中的所有文本"""
    text = ""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        print(f"  [警告] 读取 {pdf_path} 失败: {e}", file=sys.stderr)
    return text


def classify_invoice(text: str) -> str:
    """根据文本内容，判断发票类型"""
    scores = {cat: 0 for cat in CATEGORY_KEYWORDS}
    for cat, keywords in CATEGORY_KEYWORDS.items():
        for kw in keywords:
            if kw in text:
                scores[cat] += 1
    best = max(scores, key=scores.get)
    if scores[best] == 0:
        return "未分类"
    return best


def extract_amount(text: str) -> float:
    """提取发票金额（优先提取价税合计/总金额）"""
    # 优先匹配价税合计
    priority_patterns = [
        r"价税合计[（(小写）)]*\s*[¥￥]?\s*(\d+\.?\d*)",
        r"合计金额\s*[¥￥]?\s*(\d+\.?\d*)",
        r"合计\s*[¥￥]?\s*(\d+\.?\d*)",
        r"总金额\s*[¥￥]?\s*(\d+\.?\d*)",
        r"实收金额\s*[¥￥]?\s*(\d+\.?\d*)",
        r"应收金额\s*[¥￥]?\s*(\d+\.?\d*)",
        r"票价\s*[¥￥]?\s*(\d+\.?\d*)",
        r"金额合计\s*[¥￥]?\s*(\d+\.?\d*)",
        r"消费金额\s*[¥￥]?\s*(\d+\.?\d*)",
        r"行程费用\s*[¥￥]?\s*(\d+\.?\d*)",
        r"[¥￥]\s*(\d+\.\d{2})",
    ]
    for pattern in priority_patterns:
        matches = re.findall(pattern, text)
        if matches:
            # 取最大金额（通常是价税合计）
            amounts = [float(m) for m in matches if float(m) > 0]
            if amounts:
                return max(amounts)

    # 兜底：提取所有金额取最大值
    all_amounts = re.findall(r"(\d{1,6}\.\d{2})", text)
    if all_amounts:
        return max(float(a) for a in all_amounts)

    return 0.0


def extract_date(text: str) -> str:
    """提取日期"""
    patterns = [
        r"(\d{4})[年\-/](\d{1,2})[月\-/](\d{1,2})",
        r"(\d{4})(\d{2})(\d{2})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            y, m, d = match.group(1), match.group(2), match.group(3)
            return f"{y}-{int(m):02d}-{int(d):02d}"
    return ""


def extract_city_from_text(text: str) -> str:
    """从票面文本中尝试提取目的城市"""
    for city in sorted(CITY_LEVEL.keys(), key=len, reverse=True):
        if city in text:
            return city
    return ""


def process_pdf(pdf_path: str) -> dict:
    """处理单个 PDF，返回解析结果"""
    print(f"  处理: {Path(pdf_path).name}")
    text = extract_text_from_pdf(pdf_path)
    if not text.strip():
        print(f"    [警告] PDF 无文本内容（可能是图片型 PDF，请先 OCR 处理）")
        return {
            "file": Path(pdf_path).name,
            "category": "无法识别",
            "amount": 0.0,
            "date": "",
            "city": "",
            "note": "图片型PDF，需手动处理"
        }

    category = classify_invoice(text)
    amount = extract_amount(text)
    date = extract_date(text)
    city = extract_city_from_text(text)

    return {
        "file": Path(pdf_path).name,
        "category": category,
        "amount": amount,
        "date": date,
        "city": city,
        "note": ""
    }

# ─────────────────────────────────────────────
# 差旅补助计算
# ─────────────────────────────────────────────

def get_city_level(city: str) -> str:
    """查询城市等级，未找到则返回 D（其他二线城市）"""
    for key, level in CITY_LEVEL.items():
        if city in key or key in city:
            return level
    return "D"


def calc_allowance(city_days: list) -> list:
    """
    计算差旅补助
    city_days: [{"city": "北京", "days": 2}, ...]
    返回: [{"city": ..., "level": ..., "days": ..., "rate": ..., "amount": ...}]
    """
    results = []
    for item in city_days:
        city = item["city"]
        days = item["days"]
        level = get_city_level(city)
        rate = ALLOWANCE.get(level, 150)
        amount = rate * days
        results.append({
            "city": city,
            "level": f"{level}类城市",
            "days": days,
            "rate": rate,
            "amount": amount
        })
    return results

# ─────────────────────────────────────────────
# Excel 生成
# ─────────────────────────────────────────────

HEADER_COLOR = "2B5F8E"   # 深蓝色表头
ALT_ROW_COLOR = "EEF3FA"  # 交替行浅蓝
SUBTOTAL_COLOR = "D6E4F0" # 小计行颜色
TOTAL_COLOR = "C6EFCE"    # 合计行浅绿


def set_border(cell, sides="all"):
    thin = Side(style="thin")
    if sides == "all":
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    elif sides == "bottom":
        cell.border = Border(bottom=thin)


def write_excel(invoices: list, allowance_rows: list, output_path: str):
    wb = Workbook()

    # ──── Sheet 1: 报销汇总 ────
    ws1 = wb.active
    ws1.title = "报销汇总"

    title_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_COLOR)
    normal_font = Font(name="微软雅黑", size=10)
    bold_font = Font(name="微软雅黑", bold=True, size=10)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # 大标题
    ws1.merge_cells("A1:F1")
    ws1["A1"] = "出差报销汇总表"
    ws1["A1"].font = Font(name="微软雅黑", bold=True, size=14)
    ws1["A1"].alignment = center
    ws1.row_dimensions[1].height = 32

    # 生成日期
    ws1.merge_cells("A2:F2")
    ws1["A2"] = f"生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}"
    ws1["A2"].font = Font(name="微软雅黑", size=9, color="888888")
    ws1["A2"].alignment = right

    # 表头
    headers = ["序号", "类别", "发票文件", "金额（元）", "日期", "备注"]
    col_widths = [6, 12, 35, 14, 14, 20]
    ws1.row_dimensions[3].height = 24
    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws1.cell(row=3, column=i, value=h)
        cell.font = title_font
        cell.fill = header_fill
        cell.alignment = center
        set_border(cell)
        ws1.column_dimensions[get_column_letter(i)].width = w

    # 分类顺序
    categories = ["车票", "住宿费", "打车费", "未分类"]
    row_num = 4
    grand_total = 0.0
    seq = 1

    for cat in categories:
        cat_items = [inv for inv in invoices if inv["category"] == cat]
        if not cat_items:
            continue
        cat_total = 0.0
        for idx, inv in enumerate(cat_items):
            row_fill = PatternFill(fill_type="solid", fgColor="FFFFFF" if idx % 2 == 0 else ALT_ROW_COLOR)
            row_data = [seq, cat, inv["file"], inv["amount"], inv["date"], inv["note"]]
            ws1.row_dimensions[row_num].height = 20
            for col, val in enumerate(row_data, start=1):
                cell = ws1.cell(row=row_num, column=col, value=val)
                cell.font = normal_font
                cell.fill = row_fill
                cell.alignment = center if col in (1, 2, 5) else (right if col == 4 else left_align)
                set_border(cell)
            cat_total += inv["amount"]
            grand_total += inv["amount"]
            seq += 1
            row_num += 1

        # 分类小计
        ws1.row_dimensions[row_num].height = 20
        subtotal_fill = PatternFill(fill_type="solid", fgColor=SUBTOTAL_COLOR)
        subtotal_data = ["", f"{cat} 小计", "", cat_total, "", ""]
        for col, val in enumerate(subtotal_data, start=1):
            cell = ws1.cell(row=row_num, column=col, value=val)
            cell.font = bold_font
            cell.fill = subtotal_fill
            cell.alignment = center if col in (1, 5) else (right if col == 4 else center)
            set_border(cell)
        row_num += 1

    # 差旅补助行
    ws1.row_dimensions[row_num].height = 20
    allowance_total = sum(r["amount"] for r in allowance_rows)
    allowance_desc = "、".join([f"{r['city']} {r['days']}天" for r in allowance_rows])
    row_data = [seq, "差旅补助", allowance_desc, allowance_total, "", "定额补助（不报销票据）"]
    allow_fill = PatternFill(fill_type="solid", fgColor="FFFDE7")
    for col, val in enumerate(row_data, start=1):
        cell = ws1.cell(row=row_num, column=col, value=val)
        cell.font = normal_font
        cell.fill = allow_fill
        cell.alignment = center if col in (1, 2, 5) else (right if col == 4 else left_align)
        set_border(cell)
    grand_total += allowance_total
    row_num += 1

    # 总计行
    ws1.row_dimensions[row_num].height = 24
    total_fill = PatternFill(fill_type="solid", fgColor=TOTAL_COLOR)
    ws1.merge_cells(f"A{row_num}:C{row_num}")
    ws1[f"A{row_num}"] = "合计报销金额"
    ws1[f"A{row_num}"].font = Font(name="微软雅黑", bold=True, size=11)
    ws1[f"A{row_num}"].fill = total_fill
    ws1[f"A{row_num}"].alignment = center
    ws1[f"D{row_num}"] = grand_total
    ws1[f"D{row_num}"].font = Font(name="微软雅黑", bold=True, size=12, color="C00000")
    ws1[f"D{row_num}"].fill = total_fill
    ws1[f"D{row_num}"].alignment = center
    ws1[f"E{row_num}"].fill = total_fill
    ws1[f"F{row_num}"].fill = total_fill
    for col in range(1, 7):
        set_border(ws1.cell(row=row_num, column=col))

    # ──── Sheet 2: 差旅补助明细 ────
    ws2 = wb.create_sheet("差旅补助明细")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 10
    ws2.column_dimensions["D"].width = 14
    ws2.column_dimensions["E"].width = 12

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "差旅补助计算明细"
    ws2["A1"].font = Font(name="微软雅黑", bold=True, size=13)
    ws2["A1"].alignment = center
    ws2.row_dimensions[1].height = 28

    headers2 = ["出差城市", "城市等级", "天数（天）", "日补助标准（元）", "补助金额（元）"]
    ws2.row_dimensions[2].height = 22
    for i, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=2, column=i, value=h)
        cell.font = title_font
        cell.fill = header_fill
        cell.alignment = center
        set_border(cell)

    for ridx, row in enumerate(allowance_rows, start=3):
        row_fill = PatternFill(fill_type="solid", fgColor="FFFFFF" if ridx % 2 == 1 else ALT_ROW_COLOR)
        ws2.row_dimensions[ridx].height = 20
        for col, val in enumerate([row["city"], row["level"], row["days"], row["rate"], row["amount"]], start=1):
            cell = ws2.cell(row=ridx, column=col, value=val)
            cell.font = normal_font
            cell.fill = row_fill
            cell.alignment = center
            set_border(cell)

    # 小计
    total_row = len(allowance_rows) + 3
    ws2.merge_cells(f"A{total_row}:D{total_row}")
    ws2[f"A{total_row}"] = "合计"
    ws2[f"A{total_row}"].font = bold_font
    ws2[f"A{total_row}"].fill = PatternFill(fill_type="solid", fgColor=TOTAL_COLOR)
    ws2[f"A{total_row}"].alignment = center
    ws2[f"E{total_row}"] = sum(r["amount"] for r in allowance_rows)
    ws2[f"E{total_row}"].font = Font(name="微软雅黑", bold=True, size=11, color="C00000")
    ws2[f"E{total_row}"].fill = PatternFill(fill_type="solid", fgColor=TOTAL_COLOR)
    ws2[f"E{total_row}"].alignment = center
    for col in range(1, 6):
        set_border(ws2.cell(row=total_row, column=col))

    ws2.row_dimensions[total_row].height = 22

    # 备注
    ws2[f"A{total_row+2}"] = "注：出差补助为定额补助，出差期间的伙食费、市内交通费不再单独报销。"
    ws2[f"A{total_row+2}"].font = Font(name="微软雅黑", size=9, color="888888")
    ws2.merge_cells(f"A{total_row+2}:E{total_row+2}")

    wb.save(output_path)
    print(f"\n✅ Excel 报销汇总已生成: {output_path}")
    return grand_total

# ─────────────────────────────────────────────
# 主程序
# ─────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="发票报销处理脚本")
    parser.add_argument("--pdfs", nargs="+", required=True, help="PDF 发票文件路径列表")
    parser.add_argument("--city", type=str, default="", help="出差城市（逗号分隔多个城市）")
    parser.add_argument("--days", type=str, default="", help="对应天数（逗号分隔，与城市一一对应）")
    parser.add_argument("--output", type=str, default="报销汇总.xlsx", help="输出 Excel 文件路径")
    args = parser.parse_args()

    print("=" * 50)
    print("  发票报销处理中...")
    print("=" * 50)

    # 处理 PDF
    invoices = []
    for pdf_path in args.pdfs:
        result = process_pdf(pdf_path)
        invoices.append(result)

    # 展示识别结果
    print("\n📋 发票识别结果：")
    print(f"{'文件名':<30} {'类别':<10} {'金额':>10} {'日期':<12}")
    print("-" * 65)
    for inv in invoices:
        print(f"{inv['file']:<30} {inv['category']:<10} {inv['amount']:>10.2f} {inv['date']:<12}")

    # 计算差旅补助
    city_days = []
    if args.city:
        cities = [c.strip() for c in args.city.split(",") if c.strip()]
        if args.days:
            days_list = [float(d.strip()) for d in args.days.split(",") if d.strip()]
        else:
            days_list = [1.0] * len(cities)
        # 补充天数
        while len(days_list) < len(cities):
            days_list.append(1.0)
        for city, day in zip(cities, days_list):
            city_days.append({"city": city, "days": day})

    allowance_rows = calc_allowance(city_days)

    if allowance_rows:
        print("\n🏙️ 差旅补助计算：")
        print(f"{'城市':<10} {'等级':<8} {'天数':>6} {'日标准':>8} {'补助':>8}")
        print("-" * 45)
        for row in allowance_rows:
            print(f"{row['city']:<10} {row['level']:<8} {row['days']:>6} {row['rate']:>8} {row['amount']:>8.1f}")

    # 生成 Excel
    total = write_excel(invoices, allowance_rows, args.output)
    print(f"\n💰 合计报销金额：¥{total:.2f}")
    print("=" * 50)


if __name__ == "__main__":
    main()
